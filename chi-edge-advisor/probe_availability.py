#!/usr/bin/env python3
"""Live Chameleon availability with reservation windows, per node, per site.

  python probe_availability.py --rc "~/Downloads/app-cred-*-openrc.sh"
  python probe_availability.py --rc "..." --nodes            # per-node detail
  python probe_availability.py --rc "..." --xlsx live.xlsx   # advisor export
  python probe_availability.py --rc "..." --json
  python probe_availability.py --reference                   # static API check

- Live state is in Blazar (one per site); the portal Host Calendar is a Blazar view.
- Chain: Keystone /auth/tokens -> catalog -> blazar -> {resource} + /allocations.
- Resource is os-hosts everywhere except CHI@Edge, which uses devices.
- Type field is node_type on hosts, machine_name on devices.
- Three states: down (reservable=False), busy (a window covers now), else free.
- free  -> available_hours until the next reservation starts (None = unbounded).
  busy  -> next_free at the end of the contiguous booked block, not just the
           current reservation: back-to-back leases are chained.
- Terminal rounds to hours and caps at >7d. The XLSX/JSON carry exact UTC
  timestamps, which is what the advisor needs to schedule a future booking.
- An app credential is scoped to ONE site, so pass one openrc per site.
- --rc sources each openrc in a subshell, stdin closed. Secrets are never printed.
"""
from __future__ import annotations

import argparse
import glob
import json
import os
import sys
from collections import defaultdict
from datetime import datetime, timedelta, timezone


sys.path.insert(0, __file__.rsplit("/", 1)[0])

from advisor.availability.blazar import (  # noqa: E402
    CONTIGUOUS_GAP, RESOURCES, fetch_sites, load_rc, probe_site, windows,
)

HORIZON = timedelta(days=7)   # display cap; stored timestamps stay exact


def _now():
    return datetime.now(timezone.utc)


def human(delta):
    """Hours-precise, capped at the horizon. None means unbounded."""
    if delta is None or delta >= HORIZON:
        return ">7d"
    h = int(delta.total_seconds() // 3600)
    if h < 1:
        return f"{int(delta.total_seconds() // 60)}m"
    return f"{h}h" if h < 48 else f"{h // 24}d {h % 24}h"


def by_type(nodes):
    """Aggregate per node_type.

    longest_hours is the best window you could get by picking the best free
    node; None means at least one free node has nothing booked after it, so the
    window is open-ended. shortest_hours is the same set's worst case, which is
    what warns you that "free" does not mean "free for long".
    """
    groups = defaultdict(list)
    for n in nodes:
        groups[n["node_type"]].append(n)

    agg = {}
    for t, ns in sorted(groups.items()):
        free = [n for n in ns if n["status"] == "free"]
        hours = [n["available_hours"] for n in free]
        finite = [h for h in hours if h is not None]
        busy = [n["busy_hours_remaining"] for n in ns
                if n["status"] == "busy" and n["busy_hours_remaining"] is not None]
        agg[t] = {
            "free": len(free),
            "busy": sum(n["status"] == "busy" for n in ns),
            "down": sum(n["status"] == "down" for n in ns),
            "total": len(ns),
            # any unbounded free node -> best case is open-ended
            "longest_hours": None if (free and len(finite) < len(hours))
                             else (max(finite) if finite else None),
            "shortest_hours": min(finite) if finite else None,
            "soonest_free_hours": min(busy) if busy else None,
        }
    return agg


def _hrs(h):
    return human(None if h is None else timedelta(hours=h))


def render(sites, show_nodes):
    print(f"Chameleon live availability  {_now().isoformat(timespec='seconds')}")
    live = [s for s in sites if s["error"] is None]
    for s in sorted(sites, key=lambda x: x["site"]):
        print(f"\n{s['site']}")
        if s["error"]:
            print(f"  NO DATA: {s['error']}")
            continue
        n = s["nodes"]
        f = sum(x["status"] == "free" for x in n)
        d = sum(x["status"] == "down" for x in n)
        print(f"  {s['blazar']} ({s['kind']})")
        print(f"  {f} free / {len(n)} total" + (f", {d} down" if d else ""))
        print(f"    {'node_type':30} {'free/total':>11}  {'now?':4} {'window':>16}")
        for t, a in by_type(n).items():
            if a["free"]:
                now_s, win = "YES", f"up to {_hrs(a['longest_hours'])}"
                # Warn when the shortest free window is materially worse than
                # the best: "3 free" can still mean "2 of them expire tonight".
                s = a["shortest_hours"]
                if s is not None and _hrs(s) != _hrs(a["longest_hours"]):
                    win += f" (min {_hrs(s)})"
            elif a["down"] == a["total"]:
                now_s, win = "DOWN", "unreservable"
            elif a["soonest_free_hours"] is not None:
                now_s, win = "no", f"free in {_hrs(a['soonest_free_hours'])}"
            else:
                now_s, win = "no", "unknown"
            cnt = f"{a['free']}/{a['total']}" + (f" ({a['down']}d)" if a["down"] else "")
            print(f"    {t:30} {cnt:>11}  {now_s:4} {win:>16}")
        if show_nodes:
            print()
            for x in sorted(n, key=lambda x: (x["node_type"], x["name"])):
                w = (f"free {_hrs(x['available_hours'])}" if x["status"] == "free"
                     else f"busy, free in {_hrs(x['busy_hours_remaining'])}"
                     if x["status"] == "busy" else "DOWN")
                print(f"      {x['node_type']:26} {str(x['name'])[:26]:26} {w}")
    if live:
        alln = [x for s in live for x in s["nodes"]]
        print(f"\nTOTAL {sum(x['status'] == 'free' for x in alln)} free / {len(alln)} "
              f"nodes across {len(live)} sites, "
              f"{sum(x['status'] == 'down' for x in alln)} down")


def to_xlsx(sites, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
    fill = {"free": "C6EFCE", "busy": "FFEB9C", "down": "F8CBAD"}

    def sheet(ws, cols, rows, widths=None):
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            ws.cell(row=1, column=c).font = hf
            ws.cell(row=1, column=c).fill = hfill
        for r in rows:
            ws.append([r.get(c) for c in cols])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = \
                (widths or {}).get(c, min(max(12, len(c) + 2), 34))

    nodes = [n for s in sites if s["error"] is None for n in s["nodes"]]
    cols = ["site", "node_type", "name", "uid", "status", "reservable",
            "immediately_available", "available_hours", "free_until_utc",
            "busy_hours_remaining", "next_free_utc", "lease_id",
            "gpu_model", "gpu_count", "vcpus", "ram_gb"]
    ws = wb.active
    ws.title = "Nodes"
    sheet(ws, cols, nodes, {"name": 28, "uid": 38, "lease_id": 38,
                            "free_until_utc": 26, "next_free_utc": 26})
    for i, n in enumerate(nodes, start=2):
        ws.cell(row=i, column=5).fill = PatternFill(
            "solid", fgColor=fill.get(n["status"], "F2F2F2"))

    agg = []
    for s in sites:
        if s["error"]:
            continue
        for t, a in by_type(s["nodes"]).items():
            agg.append({"site": s["site"], "node_type": t, **a,
                        "longest_window": _hrs(a["longest_hours"]) if a["free"] else "",
                        "shortest_window": _hrs(a["shortest_hours"])
                        if a["shortest_hours"] is not None else "",
                        "soonest_free": _hrs(a["soonest_free_hours"])
                        if a["soonest_free_hours"] is not None else ""})
    sheet(wb.create_sheet("By node_type"),
          ["site", "node_type", "free", "busy", "down", "total",
           "longest_window", "longest_hours", "shortest_window", "shortest_hours",
           "soonest_free", "soonest_free_hours"],
          agg)

    sheet(wb.create_sheet("Sites"),
          ["site", "blazar", "kind", "free", "busy", "down", "total", "error"],
          [{"site": s["site"], "blazar": s["blazar"], "kind": s["kind"],
            "free": sum(x["status"] == "free" for x in s["nodes"]),
            "busy": sum(x["status"] == "busy" for x in s["nodes"]),
            "down": sum(x["status"] == "down" for x in s["nodes"]),
            "total": len(s["nodes"]), "error": s["error"]} for s in sites],
          {"blazar": 46, "error": 40})

    wb.save(path)
    print(f"\n[written] {path}  ({len(nodes)} nodes, exact UTC timestamps)")


def reference_check():
    """Does the public reference API expose live state? Settled: it does not."""
    from advisor.availability.reference_api import ReferenceApiBackend
    from advisor.config import settings
    base, site = settings.reference_api_base.rstrip("/"), settings.edge_site_uid
    b = ReferenceApiBackend(base=base, site=site)
    live = b.probe_live_status()
    return {"base": base, "edge_devices_exposed": len(b.list_devices()),
            "live_endpoint": live["live_endpoint"],
            "verdict": "LIVE-AWARE" if live["live_endpoint"] else "STATIC-ONLY"}


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--rc", help="glob of openrc files, one per site")
    ap.add_argument("--nodes", action="store_true", help="per-node detail")
    ap.add_argument("--xlsx", help="write the advisor export to this path")
    ap.add_argument("--reference", action="store_true",
                    help="also run the static reference-API check")
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args()

    ref = reference_check() if args.reference else None
    sites = []
    if args.rc:
        for p in sorted(glob.glob(os.path.expanduser(args.rc))):
            sites.append(probe_site(load_rc(p), p))
    elif not args.reference:
        ap.error("--rc is required for live state (or use --reference)")

    if args.json:
        print(json.dumps({"generated_utc": _now().isoformat(),
                          "reference": ref, "sites": sites}, indent=2))
    else:
        if ref:
            print(f"reference API {ref['base']}: {ref['verdict']}, "
                  f"{ref['edge_devices_exposed']} edge devices, "
                  f"live endpoint {ref['live_endpoint']}\n")
        if sites:
            render(sites, args.nodes)
    if args.xlsx and sites:
        to_xlsx(sites, args.xlsx)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
