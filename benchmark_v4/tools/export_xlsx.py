"""Export the item bank + grading workspace to exports/benchmark_v4.xlsx."""
import json
from pathlib import Path
import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent.parent
HDR = Font(bold=True, color="FFFFFF", name="Arial")
FILL = PatternFill("solid", start_color="1F6F5C")
BODY = Font(name="Arial")

tiers = {}
for row in json.load(open(ROOT / "exports" / "tier_report.json"))["rows"]:
    tiers[(row["item"], row["condition"])] = row["tier"]

items = [yaml.safe_load(p.read_text()) for p in sorted((ROOT / "items").glob("*.yaml"))]

wb = Workbook()

ws = wb.active
ws.title = "Items"
cols = ["ID", "Lineage", "Category", "Axis", "Tier (matched)", "Verification",
        "Prompt (paste verbatim)", "Designed trap", "Trap tags", "Target artifact",
        "Fed (matched)", "Fed (heldout)", "# checkers"]
ws.append(cols)
for i, it in enumerate(items, start=2):
    fs = it.get("fed_sets", {})
    matched_cond = "matched" if "matched" in fs else ("uncovered" if "uncovered" in fs else "")
    ws.append([
        it["id"], it.get("lineage", ""), it["category"], it["axis"],
        tiers.get((it["id"], matched_cond), ""), it["verification_level"],
        it["prompt"], it.get("designed_trap", ""),
        ", ".join(it.get("trap_tags", [])), ", ".join(it.get("target_artifact", [])),
        ", ".join(fs.get("matched", fs.get("uncovered", []))),
        ", ".join(fs.get("heldout", [])), len(it.get("checkers", [])),
    ])
for c in ws[1]:
    c.font, c.fill = HDR, FILL
for col, w in zip("ABCDEFGHIJKLM", [7, 9, 20, 9, 12, 11, 70, 45, 28, 12, 12, 12, 10]):
    ws.column_dimensions[col].width = w
for r in ws.iter_rows(min_row=2):
    for c in r:
        c.font = BODY
        c.alignment = Alignment(vertical="top", wrap_text=c.column_letter in "GH")

ws2 = wb.create_sheet("Output")
ws2.append(["ID", "Condition", "System", "Model output (paste here)",
            "Checker verdict (harness)", "Human 0-2", "Failure tag", "Notes"])
for c in ws2[1]:
    c.font, c.fill = HDR, FILL
for col, w in zip("ABCDEFGH", [7, 11, 18, 70, 20, 10, 16, 40]):
    ws2.column_dimensions[col].width = w

ws3 = wb.create_sheet("RunMatrix")
ws3.append(["ID", "Condition", "Fed artifacts", "Computed tier"])
for row in json.load(open(ROOT / "exports" / "tier_report.json"))["rows"]:
    ws3.append([row["item"], row["condition"], ", ".join(row["fed"]), row["tier"]])
for c in ws3[1]:
    c.font, c.fill = HDR, FILL
for col, w in zip("ABCD", [7, 11, 20, 13]):
    ws3.column_dimensions[col].width = w

ws4 = wb.create_sheet("Summary")
ws4["B2"] = "Results summary (auto-computes as the Output sheet is graded)"
ws4["B2"].font = Font(bold=True, name="Arial", size=12)
ws4.append([])
ws4.append(["", "Condition", "Graded cells", "Checker pass rate", "Mean human 0-2"])
for c in ws4[4][1:]:
    c.font, c.fill = HDR, FILL
for i, cond in enumerate(["blind", "matched", "heldout", "uncovered"], start=5):
    ws4[f"B{i}"] = cond
    ws4[f"C{i}"] = f'=COUNTIFS(Output!B:B,"{cond}",Output!F:F,"<>")'
    ws4[f"D{i}"] = (f'=IFERROR(COUNTIFS(Output!B:B,"{cond}",Output!E:E,"PASS")/'
                    f'COUNTIFS(Output!B:B,"{cond}",Output!E:E,"<>"),"")')
    ws4[f"E{i}"] = f'=IFERROR(AVERAGEIFS(Output!F:F,Output!B:B,"{cond}"),"")'
for col, w in zip("BCDE", [12, 12, 16, 14]):
    ws4.column_dimensions[col].width = w

wb.save(ROOT / "exports" / "benchmark_v4.xlsx")
print("wrote exports/benchmark_v4.xlsx")
